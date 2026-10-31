from contextlib import contextmanager
from pathlib import Path
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from matplotlib.ticker import FuncFormatter


FOLDER = Path(__file__).parent
OUTPUT_FOLDER = FOLDER / "eredmenyek"
CORE_PROPERTIES_NS = "{http://purl.org/dc/terms/}"

SPORTAGAK = [
	("Amerikai futball", ("FOOTBALL",)),
	("Baseball", ("BASEBALL",)),
	("Softball", ("SOFTBALL",)),
	("Kosárlabda", ("BASKETBALL",)),
	("Labdarúgás", ("SOCCER",)),
	("Kézilabda", ("HANDBALL", "HAND BALL")),
	("Röplabda", ("VOLLEYBALL",)),
	("Hoki", ("HOCKEY",)),
	("Asztalitenisz", ("TABLE TENNIS",)),
	("Tenisz", ("TENNIS",)),
	("Golf", ("GOLF",)),
	("Gimnasztika", ("GYMNASTICS",)),
	("Birkózás", ("WRESTLING",)),
	("Lacrosse", ("LACROSSE",)),
	("Rögbi", ("RUGBY",)),
	("Cheerleading", ("CHEERLEADING",)),
	("Úszás", ("SWIMMING",)),
	("Kerékpározás", ("BICYCLE", "CYCLING")),
	("Gördeszkázás", ("SKATEBOARD",)),
	("Síelés", ("SKIING", "SKI EQUIPMENT")),
	("Snowboard", ("SNOWBOARD",)),
]


def egysegesites(value):
	if pd.isna(value):
		return ""
	try:
		return str(int(float(value)))
	except (TypeError, ValueError):
		return str(value).strip()


@contextmanager
def neiisadatokmegnyit(file_path):
	javitott_adatok = None
	try:
		adatok = load_workbook(file_path, read_only=True, data_only=True)
	except ValueError as error:
		if "could not read properties" not in str(error):
			raise

		javitott_adatok = BytesIO()
		with ZipFile(file_path) as ahelye, ZipFile(
			javitott_adatok, "w", compression=ZIP_DEFLATED
		) as javitottremelem:
			for cucc in ahelye.infolist():
				acucc = ahelye.read(cucc.filename)
				if cucc.filename == "docProps/core.xml":
					root = ElementTree.fromstring(acucc)
					for name in ("created", "modified"):
						datuma = root.find(CORE_PROPERTIES_NS + name)
						if datuma is not None:
							root.remove(datuma)
					acucc = ElementTree.tostring(
						root, encoding="utf-8", xml_declaration=True
					)
				javitottremelem.writestr(cucc, acucc)
		javitott_adatok.seek(0)
		try:
			adatok = load_workbook(javitott_adatok, read_only=True, data_only=True)
		except Exception:
			javitott_adatok.close()
			raise

	try:
		yield adatok
	finally:
		adatok.close()
		if javitott_adatok is not None:
			javitott_adatok.close()


def beolvasas(file_path):
	with neiisadatokmegnyit(file_path) as workbook:
		sheet = workbook["NEISS_FMT"]
		sheet.reset_dimensions()
		rows = sheet.iter_rows(values_only=True)
		next(rows, None)
		product_labels = {}
		diagnosis_labels = {}
		body_part_labels = {}
		for row in rows:
			if row[0] in ("PROD", "DIAG", "BDYPT") and row[1] is not None and row[3] is not None:
				labels = {
					"PROD": product_labels,
					"DIAG": diagnosis_labels,
					"BDYPT": body_part_labels,
				}[row[0]]
				labels[egysegesites(row[1])] = str(row[3]).strip().upper()
		return product_labels, diagnosis_labels, body_part_labels


def reszletes_serulestipus(narrative, body_part_label, diagnosis_code):
	narrative = str(narrative or "").upper()
	specifikus_serulesek = (
		(r"\bACL\b|ANTERIOR CRUCIATE", "ACL-sérülés (elülső keresztszalag)"),
		(r"\bMCL\b|MEDIAL COLLATERAL", "MCL-sérülés (belső oldalszalag)"),
		(r"\bPCL\b|POSTERIOR CRUCIATE", "PCL-sérülés (hátsó keresztszalag)"),
		(r"\bLCL\b|LATERAL COLLATERAL", "LCL-sérülés (külső oldalszalag)"),
		(r"\bATFL\b|ANTERIOR TALOFIBULAR", "ATFL-sérülés (bokaszalag)"),
		(r"\bCFL\b|CALCANEOFIBULAR", "CFL-sérülés (bokaszalag)"),
		(r"\bMENISC(?:US|AL)\b", "Meniszkuszsérülés"),
		(r"\bACHILLES\b", "Achilles-ín sérülés"),
		(r"ROTATOR CUFF", "Rotátorköpeny-sérülés"),
	)
	for pattern, label in specifikus_serulesek:
		if re.search(pattern, narrative):
			return label

	diagnosis = egysegesites(diagnosis_code)
	testreszek = {
		"SHOULDER": "Váll",
		"WRIST": "Csukló",
		"KNEE": "Térd",
		"ANKLE": "Boka",
		"FOOT": "Lábfej",
	}
	serulesek = {
		"53": "zúzódás / horzsolás",
		"55": "ficam",
		"57": "törés",
		"59": "vágott seb",
		"64": "rándulás / húzódás",
	}
	body_part_label = str(body_part_label or "").upper()
	for body_part_name, testresz in testreszek.items():
		if body_part_name in body_part_label:
			serules = serulesek.get(diagnosis, "egyéb sérülés")
			return f"{testresz} – {serules}"
	return None


def sportrend(product_label): 
	if not isinstance(product_label, str):
		return None

	for sport, keywords in SPORTAGAK:
		if any(keyword in product_label for keyword in keywords):
			return sport

	if "SPORTS AND RECREATIONAL ACTIVITY" in product_label:
		return "Egyéb sport/rekreáció"
	return None


def serulesbetolt(files):
	osszadat = []
	required_columns = [
		"CPSC_Case_Number",
		"Sex",
		"Product_1",
		"Diagnosis",
		"Body_Part",
		"Narrative_1",
		"Weight",
	]
	for file_path in files:
		print(f"Feldolgozás: {file_path.stem.replace('neiss', '')}...", flush=True)
		product_labels, diagnosis_labels, body_part_labels = beolvasas(file_path)
		with neiisadatokmegnyit(file_path) as workbook:
			sheet = workbook.worksheets[0]
			sheet.reset_dimensions()
			rows = sheet.iter_rows(values_only=True)
			headers = next(rows)
			oszlop_indexes = {name: headers.index(name) for name in required_columns}
			esetek = []
			for row in rows:
				if row[oszlop_indexes["CPSC_Case_Number"]] is None:
					continue
				product_kod = egysegesites(row[oszlop_indexes["Product_1"]])
				sport = sportrend(product_labels.get(product_kod))
				diagnosis_kod = egysegesites(row[oszlop_indexes["Diagnosis"]])
				weight = row[oszlop_indexes["Weight"]]
				if sport is None or weight is None:
					continue
				try:
					weight = float(weight)
				except (TypeError, ValueError):
					continue
				nem = row[oszlop_indexes["Sex"]]
				nem = {0: "Ismeretlen", 1: "Férfi", 2: "Nő"}.get(nem, "Ismeretlen")
				serulestipus = diagnosis_labels.get(
					diagnosis_kod,
					f"ISMERETLEN ({diagnosis_kod})" if diagnosis_kod else "ISMERETLEN SÉRÜLÉSTÍPUS",
				)
				reszletes_tipus = reszletes_serulestipus(
					row[oszlop_indexes["Narrative_1"]],
					body_part_labels.get(
						egysegesites(row[oszlop_indexes["Body_Part"]]), ""
					),
					diagnosis_kod,
				)
				esetek.append(
					{
						"sport": sport,
						"nem": nem,
						"serulestipus": serulestipus,
						"testrész": body_part_labels.get(
							egysegesites(row[oszlop_indexes["Body_Part"]]),
							"ISMERETLEN TESTRÉSZ",
						),
						"reszletes_serules": reszletes_tipus,
						"Weight": weight,
					}
				)
			osszadat.extend(esetek)

	if not osszadat:
		raise ValueError("Nem találtam feldolgozható NEISS-adatot.")
	return pd.DataFrame(
		osszadat,
		columns=[
			"sport",
			"nem",
			"serulestipus",
			"testrész",
			"reszletes_serules",
			"Weight",
		],
	)


def diagment(data):
	nemenkent = data[data["nem"].isin(["Férfi", "Nő"])].pivot_table(
		index="sport", columns="nem", values="Weight", aggfunc="sum", fill_value=0
	)
	nemenkent = nemenkent.reindex(columns=["Férfi", "Nő"], fill_value=0)
	nemenkent = nemenkent.loc[nemenkent.sum(axis=1) > 0]
	nemenkent = nemenkent.loc[nemenkent.sum(axis=1).sort_values().index]
	fig, ax = plt.subplots(figsize=(13, max(7, len(nemenkent) * 0.62)))
	fig.patch.set_facecolor("white")
	ax.set_facecolor("#f7f9fa")
	nemenkent.plot(kind="barh", ax=ax, color={"Férfi": "#177e89", "Nő": "#e07a5f"})
	for bars in ax.containers:
		labels = [
			f"{bar.get_width():,.1f}".replace(",", " ").replace(".", ",")
			for bar in bars
		]
		ax.bar_label(bars, labels=labels, padding=4, fontsize=8)
	ax.set_xlim(0, nemenkent.to_numpy().max() * 1.24)
	ax.set_title("Sportserülések sportág és nem szerint (2016–2025)", fontsize=16, pad=14)
	ax.set_xlabel("Sérülések száma (millió)")
	ax.set_ylabel("")
	ax.xaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value / 1_000_000:g}"))
	ax.grid(axis="x", color="#dce3e6", linewidth=0.8)
	ax.set_axisbelow(True)
	ax.tick_params(axis="y", labelsize=9)
	ax.legend(title="Nem", frameon=False, loc="lower right")
	fig.tight_layout()
	fig.savefig(OUTPUT_FOLDER / "sportag_es_nem_szerint.png", dpi=160)
	plt.close(fig)


def serulestipusdiagram(data):
	serulesek = data.dropna(subset=["reszletes_serules"])
	if serulesek.empty:
		raise ValueError(
			"Nem találtam a leírás vagy testrész alapján részletesen besorolható sérülést."
		)
	serulestipusok = serulesek.pivot_table(
		index="sport",
		columns="reszletes_serules",
		values="Weight",
		aggfunc="sum",
		fill_value=0,
	)
	kiemelt_tipusok = [
		"ACL-sérülés (elülső keresztszalag)",
		"MCL-sérülés (belső oldalszalag)",
		"PCL-sérülés (hátsó keresztszalag)",
		"LCL-sérülés (külső oldalszalag)",
		"Meniszkuszsérülés",
		"ATFL-sérülés (bokaszalag)",
		"CFL-sérülés (bokaszalag)",
		"Achilles-ín sérülés",
		"Rotátorköpeny-sérülés",
		"Boka – rándulás / húzódás",
		"Boka – törés",
		"Boka – ficam",
		"Boka – egyéb sérülés",
		"Térd – rándulás / húzódás",
		"Térd – törés",
		"Térd – ficam",
	]
	tipus_sorrend = [
		tipus for tipus in kiemelt_tipusok if tipus in serulestipusok.columns
	]
	if not tipus_sorrend:
		raise ValueError(
			"Nem találtam ACL/MCL-, boka- vagy térdsérülésként besorolható esetet."
		)
	oszlopok = 2
	sorok = (len(tipus_sorrend) + oszlopok - 1) // oszlopok
	fig, axes = plt.subplots(
		sorok,
		oszlopok,
		figsize=(15, max(5, sorok * 3)),
		squeeze=False,
	)
	fig.patch.set_facecolor("white")
	for index, tipus in enumerate(tipus_sorrend):
		ax = axes[index // oszlopok][index % oszlopok]
		ax.set_facecolor("#f7f9fa")
		adat = serulestipusok[tipus].nlargest(5).sort_values()
		ax.barh(adat.index, adat.values, color="#177e89")
		ax.margins(x=0.2)
		ax.set_title(f"{tipus} (összesen: {serulestipusok[tipus].sum():,.0f})", fontsize=10)
		ax.set_xlabel("Becsült esetszám")
		ax.grid(axis="x", color="#dce3e6", linewidth=0.8)
		ax.set_axisbelow(True)
		ax.tick_params(axis="y", labelsize=8)
		for bar in ax.containers:
			ax.bar_label(
				bar,
				labels=[f"{ertek:,.0f}".replace(",", " ") for ertek in adat.values],
				padding=3,
				fontsize=8,
			)
	for index in range(len(tipus_sorrend), sorok * oszlopok):
		axes[index // oszlopok][index % oszlopok].set_visible(False)
	fig.text(
		0.5,
		0.995,
		"Részletes sérüléstípusok sportáganként (2016–2025)",
		ha="center",
		va="top",
		fontsize=16,
	)
	fig.text(
		0.5,
		0.977,
		"ACL/MCL/PCL/LCL, boka, térd és hasonló sérülések; panelek: 5 leggyakoribb sportág",
		ha="center",
		va="top",
		fontsize=10,
	)
	fig.tight_layout(rect=(0, 0, 1, 0.96))
	fig.savefig(OUTPUT_FOLDER / "sportag_es_serulestipus_szerint.png", dpi=160)
	plt.close(fig)


def kezilabda_diagram(data):
	kezilabda = data[data["sport"] == "Kézilabda"]
	if kezilabda.empty:
		raise ValueError("Nem találtam kézilabdához besorolt esetet.")

	serulestipusok = kezilabda.groupby("serulestipus")["Weight"].sum().nlargest(12)
	testrészek = kezilabda.groupby("testrész")["Weight"].sum().nlargest(10)
	serulesnevek = {
		"STRAIN, SPRAIN": "Rándulás / húzódás",
		"FRACTURE": "Törés",
		"CONTUSIONS, ABR.": "Zúzódás / horzsolás",
		"LACERATION": "Vágott seb",
		"CONCUSSION": "Agyrázkódás",
		"DISLOCATION": "Ficam",
		"INTERNAL INJURY": "Belső sérülés",
		"OTHER": "Egyéb",
		"HEMATOMA": "Vérömleny",
		"NERVE DAMAGE": "Idegkárosodás",
		"POISONING": "Mérgezés",
		"DERMA/CONJUNCT": "Bőr-/kötőhártya-sérülés",
	}
	testrésznevek = {
		"KNEE": "Térd",
		"ANKLE": "Boka",
		"FINGER": "Ujj",
		"FOOT": "Lábfej",
		"HAND": "Kéz",
		"WRIST": "Csukló",
		"LOWER LEG": "Lábszár",
		"HEAD": "Fej",
		"FACE": "Arc",
		"LOWER TRUNK": "Alsó törzs",
		"UPPER TRUNK": "Felső törzs",
		"UPPER TRUNK(OLD)": "Felső törzs",
		"SHOULDER": "Váll",
		"UPPER LEG": "Comb",
	}
	serulestipusok.index = [
		serulesnevek.get(tipus.split(" - ", 1)[-1], tipus.split(" - ", 1)[-1].title())
		for tipus in serulestipusok.index
	]
	testrészek.index = [
		testrésznevek.get(testresz.split(" - ", 1)[-1], testresz.split(" - ", 1)[-1].title())
		for testresz in testrészek.index
	]

	fig, axes = plt.subplots(1, 2, figsize=(15, 8))
	fig.patch.set_facecolor("white")
	for ax, values, title, color in (
		(axes[0], serulestipusok, "Sérüléstípusok", "#177e89"),
		(axes[1], testrészek, "Érintett testrészek", "#e07a5f"),
	):
		values = values.sort_values()
		ax.set_facecolor("#f7f9fa")
		bars = ax.barh(values.index, values.values, color=color)
		ax.set_title(title, fontsize=13)
		ax.set_xlabel("Országos becsült esetszám")
		ax.grid(axis="x", color="#dce3e6", linewidth=0.8)
		ax.set_axisbelow(True)
		ax.tick_params(axis="y", labelsize=9)
		ax.margins(x=0.2)
		ax.bar_label(
			bars,
			labels=[f"{ertek:,.0f}".replace(",", " ") for ertek in values.values],
			padding=3,
			fontsize=8,
		)
	fig.suptitle("Kézilabda-sérülések (2016–2025)", fontsize=16)
	fig.tight_layout(rect=(0, 0, 1, 0.95))
	fig.savefig(OUTPUT_FOLDER / "kezilabda_serulesek.png", dpi=160)
	plt.close(fig)

	serulesenkent = (
		kezilabda.groupby("serulestipus", as_index=False)["Weight"]
		.sum()
		.assign(Bontás="Sérüléstípus")
		.rename(columns={"serulestipus": "Kategória", "Weight": "Becsült esetszám"})
	)
	testreszenkent = (
		kezilabda.groupby("testrész", as_index=False)["Weight"]
		.sum()
		.assign(Bontás="Testrész")
		.rename(columns={"testrész": "Kategória", "Weight": "Becsült esetszám"})
	)
	pd.concat([serulesenkent, testreszenkent], ignore_index=True).to_csv(
		OUTPUT_FOLDER / "kezilabda_serulesek.csv",
		index=False,
		encoding="utf-8-sig",
	)


def main():
	files = sorted(FOLDER.glob("neiss20*.xlsx"))
	if not files:
		raise FileNotFoundError(
			"Nem találtam neiss20xx.xlsx nevű fájlt. Tedd az Excel-fájlokat a valmai.py mellé."
		)

	OUTPUT_FOLDER.mkdir(exist_ok=True)
	serulesek = serulesbetolt(files)
	if serulesek.empty:
		raise ValueError(
			"Nem sikerült sportághoz sorolható esetet találni. Ellenőrizd az Excel-fájlok termékkódjait."
		)

	nemeenkent = (
		serulesek.groupby(["sport", "nem"], as_index=False)["Weight"]
		.sum()
		.rename(columns={"sport": "Sportág", "nem": "Nem", "Weight": "Becsült esetszám"})
	)
	nemeenkent.to_csv(OUTPUT_FOLDER / "sportag_es_nem.csv", index=False, encoding="utf-8-sig")
	serulestipusonkent = (
		serulesek.groupby(["sport", "serulestipus"], as_index=False)["Weight"]
		.sum()
		.rename(
			columns={
				"sport": "Sportág",
				"serulestipus": "Sérüléstípus",
				"Weight": "Becsült esetszám",
			}
		)
	)
	serulestipusonkent.to_csv(
		OUTPUT_FOLDER / "sportag_es_serulestipus.csv",
		index=False,
		encoding="utf-8-sig",
	)
	reszletes_serulesenkent = (
		serulesek.dropna(subset=["reszletes_serules"])
		.groupby(["sport", "reszletes_serules"], as_index=False)["Weight"]
		.sum()
		.rename(
			columns={
				"sport": "Sportág",
				"reszletes_serules": "Részletes sérülés",
				"Weight": "Becsült esetszám",
			}
		)
	)
	reszletes_serulesenkent.to_csv(
		OUTPUT_FOLDER / "sportag_es_reszletes_serules.csv",
		index=False,
		encoding="utf-8-sig",
	)
	diagment(serulesek)
	serulestipusdiagram(serulesek)
	kezilabda_diagram(serulesek)

	print(f"Besorolt NEISS-esetek: {len(serulesek):,}".replace(",", " "))
	print(f"Országos becslés a Weight mező alapján: {serulesek['Weight'].sum():,.0f}".replace(",", " "))
	print(f"Elkészült eredmények mappája: {OUTPUT_FOLDER}")

if __name__ == "__main__":
	main()